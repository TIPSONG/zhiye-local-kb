from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import math
import os
import re
import sqlite3
from collections import Counter
from contextlib import contextmanager
from pathlib import Path
from typing import Iterable
from urllib.parse import urlparse

import httpx
import numpy as np
from bs4 import BeautifulSoup
from docx import Document as WordDocument
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from dotenv import load_dotenv
from starlette.middleware.trustedhost import TrustedHostMiddleware
from fastapi.responses import JSONResponse
from pypdf import PdfReader
from pptx import Presentation


APP_DIR = Path(__file__).resolve().parent
load_dotenv(APP_DIR / '.env')
def local_path(value: str) -> Path:
    path = Path(value).expanduser()
    return (path if path.is_absolute() else APP_DIR / path).resolve()

DB_PATH = local_path(os.getenv('RAG_DB_PATH', 'data/knowledge.db'))
DOCUMENTS_DIR = local_path(os.getenv('RAG_DOCUMENTS_DIR', 'documents'))
LLM_URL = os.getenv("RAG_LLM_URL", "http://127.0.0.1:8080/v1").rstrip("/")
EMBED_URL = os.getenv("RAG_EMBED_URL", "http://127.0.0.1:8081/v1").rstrip("/")
RERANK_URL = os.getenv("RAG_RERANK_URL", "http://127.0.0.1:8082/v1").rstrip("/")
LLM_MODEL = os.getenv("RAG_LLM_MODEL", "local-chat")
EMBED_MODEL = os.getenv("RAG_EMBED_MODEL", "bge-m3")
RERANK_MODEL = os.getenv("RAG_RERANK_MODEL", "bge-reranker-v2-m3")
CHUNK_SIZE = int(os.getenv("RAG_CHUNK_SIZE", "1200"))
CHUNK_OVERLAP = int(os.getenv("RAG_CHUNK_OVERLAP", "200"))
RETRIEVAL_CANDIDATES = int(os.getenv("RAG_RETRIEVAL_CANDIDATES", "40"))
CONTEXT_CHUNKS = int(os.getenv("RAG_CONTEXT_CHUNKS", "8"))
REQUEST_TIMEOUT = float(os.getenv("RAG_REQUEST_TIMEOUT", "180"))
MAX_FILE_BYTES = int(os.getenv('RAG_MAX_FILE_BYTES', '52428800'))
if CHUNK_SIZE < 32 or not 0 <= CHUNK_OVERLAP < CHUNK_SIZE:
    raise ValueError('Require CHUNK_SIZE >= 32 and 0 <= CHUNK_OVERLAP < CHUNK_SIZE')

def headers(kind: str) -> dict:
    key = os.getenv(f'RAG_{kind}_API_KEY', '')
    return {'Authorization': f'Bearer {key}'} if key else {}

SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".rst", ".log", ".py", ".js", ".ts", ".tsx", ".jsx",
    ".java", ".go", ".rs", ".c", ".cpp", ".h", ".hpp", ".json", ".yaml",
    ".yml", ".toml", ".ini", ".xml", ".csv", ".html", ".htm", ".pdf",
    ".docx", ".pptx", ".xlsx",
}

app = FastAPI(title="Local RAG Knowledge Base", version="0.1.0")
app.add_middleware(TrustedHostMiddleware, allowed_hosts=['127.0.0.1','localhost','[::1]'])

@app.middleware('http')
async def local_origin(request, call_next):
    origin = request.headers.get('origin')
    if origin and origin != f'{request.url.scheme}://{request.headers.get("host", "")}':
        return JSONResponse({'detail': 'Cross-origin access is not allowed'}, status_code=403)
    return await call_next(request)
app.mount("/static", StaticFiles(directory=APP_DIR / "static"), name="static")


@app.get("/", include_in_schema=False)
def frontend():
    return FileResponse(APP_DIR / "static" / "index.html")


class IngestRequest(BaseModel):
    path: str = Field(min_length=1, max_length=4096)
    collection: str = Field(default="default", min_length=1, max_length=100)
    recursive: bool = True


class QueryRequest(BaseModel):
    question: str = Field(min_length=1, max_length=6000)
    collection: str = Field(default="default", min_length=1, max_length=100)
    top_k: int = Field(default=CONTEXT_CHUNKS, ge=1, le=20)
    generate: bool = True


@contextmanager
def connect_db():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA journal_mode=WAL")
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS chunks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            collection TEXT NOT NULL,
            source TEXT NOT NULL,
            file_hash TEXT NOT NULL,
            chunk_index INTEGER NOT NULL,
            title TEXT NOT NULL,
            content TEXT NOT NULL,
            embedding BLOB NOT NULL,
            embedding_dim INTEGER NOT NULL,
            UNIQUE(collection, source, chunk_index)
        )
        """
    )
    connection.execute("CREATE INDEX IF NOT EXISTS idx_chunks_collection ON chunks(collection)")
    try:
        with connection:
            yield connection
    finally:
        connection.close()


def read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        doc = WordDocument(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n\n".join(parts)
    if suffix == ".pptx":
        deck = Presentation(str(path))
        parts: list[str] = []
        for number, slide in enumerate(deck.slides, 1):
            text = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            parts.append(f"# Slide {number}\n" + "\n".join(text))
        return "\n\n".join(parts)
    if suffix == ".xlsx":
        workbook = load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for sheet in workbook.worksheets:
            parts.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    parts.append(" | ".join(values))
        workbook.close()
        return "\n".join(parts)
    raw = path.read_text(encoding="utf-8", errors="replace")
    if suffix in {".html", ".htm"}:
        return BeautifulSoup(raw, "html.parser").get_text("\n")
    if suffix == ".csv":
        rows = csv.reader(io.StringIO(raw))
        return "\n".join(" | ".join(row) for row in rows)
    return raw


def document_title(text: str, fallback: str) -> str:
    for line in text.splitlines():
        candidate = re.sub(r"^\s*#+\s*", "", line).strip()
        if candidate:
            return candidate[:160]
    return fallback


def split_chunks(text: str) -> list[str]:
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    chunks, start = [], 0
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        piece = text[start:end].strip()
        if piece: chunks.append(piece)
        if end == len(text): break
        start = end - CHUNK_OVERLAP
    return chunks


def embed_texts(texts: list[str], batch_size: int = 16) -> list[np.ndarray]:
    vectors: list[np.ndarray] = []
    with httpx.Client(timeout=REQUEST_TIMEOUT, headers=headers('EMBED')) as client:
        for start in range(0, len(texts), batch_size):
            batch = texts[start:start + batch_size]
            response = client.post(
                f"{EMBED_URL}/embeddings",
                json={"model": EMBED_MODEL, "input": batch},
            )
            response.raise_for_status()
            data = sorted(response.json()["data"], key=lambda item: item.get("index", 0))
            if len(data) != len(batch) or [item.get('index') for item in data] != list(range(len(batch))):
                raise ValueError('Embedding endpoint returned missing or unordered indices')
            for item in data:
                vector = np.asarray(item["embedding"], dtype=np.float32)
                norm = float(np.linalg.norm(vector))
                if vector.ndim != 1 or not vector.size or not np.isfinite(vector).all() or not norm:
                    raise ValueError('Invalid embedding vector')
                vectors.append(vector / norm if norm else vector)
    return vectors


def tokenize(text: str) -> list[str]:
    lowered = text.lower()
    latin = re.findall(r"[a-z0-9_]+", lowered)
    cjk_runs = re.findall(r"[\u3400-\u9fff]+", lowered)
    cjk: list[str] = []
    for run in cjk_runs:
        cjk.extend(run)
        cjk.extend(run[index:index + 2] for index in range(len(run) - 1))
    return latin + cjk


def bm25_scores(query: str, documents: list[str], k1: float = 1.5, b: float = 0.75) -> np.ndarray:
    tokenized = [tokenize(document) for document in documents]
    query_tokens = tokenize(query)
    count = len(tokenized)
    if not count:
        return np.zeros(0, dtype=np.float32)
    avg_len = sum(len(tokens) for tokens in tokenized) / count or 1.0
    document_frequency: Counter[str] = Counter()
    for tokens in tokenized:
        document_frequency.update(set(tokens))
    scores = np.zeros(count, dtype=np.float32)
    for index, tokens in enumerate(tokenized):
        frequencies = Counter(tokens)
        length_norm = k1 * (1 - b + b * len(tokens) / avg_len)
        for token in query_tokens:
            frequency = frequencies.get(token, 0)
            if not frequency:
                continue
            df = document_frequency[token]
            inverse_frequency = math.log(1 + (count - df + 0.5) / (df + 0.5))
            scores[index] += inverse_frequency * frequency * (k1 + 1) / (frequency + length_norm)
    return scores


def reciprocal_rank_fusion(vector_scores: np.ndarray, keyword_scores: np.ndarray, limit: int) -> list[int]:
    fused: Counter[int] = Counter()
    for scores in (vector_scores, keyword_scores):
        order = np.argsort(-scores)[:limit]
        for rank, index in enumerate(order, 1):
            fused[int(index)] += 1.0 / (60 + rank)
    return [index for index, _ in fused.most_common(limit)]


def rerank(question: str, rows: list[sqlite3.Row], top_k: int) -> list[dict]:
    payload = {
        "model": RERANK_MODEL,
        "query": question,
        "documents": [row["content"] for row in rows],
        "top_n": min(top_k, len(rows)),
    }
    with httpx.Client(timeout=REQUEST_TIMEOUT, headers=headers('RERANK')) as client:
        response = client.post(f"{RERANK_URL}/rerank", json=payload)
        response.raise_for_status()
        body = response.json()
    items = body.get("results", body.get("data", []))
    ranked = []
    seen = set()
    for item in items:
        index = int(item["index"])
        if index < 0 or index >= len(rows) or index in seen:
            raise ValueError('Invalid reranker result index')
        seen.add(index)
        score = item.get("relevance_score", item.get("score", 0.0))
        ranked.append({"row": rows[index], "score": float(score)})
    return sorted(ranked, key=lambda item: item['score'], reverse=True)[:top_k]


def retrieve(question: str, collection: str, top_k: int) -> list[dict]:
    with connect_db() as connection:
        rows = connection.execute(
            "SELECT * FROM chunks WHERE collection = ? ORDER BY id", (collection,)
        ).fetchall()
    if not rows:
        raise HTTPException(status_code=404, detail=f"Collection '{collection}' is empty")
    query_vector = embed_texts([question])[0]
    if any(row['embedding_dim'] != len(query_vector) for row in rows):
        raise HTTPException(409, 'Embedding dimension changed: use a new collection and reimport documents')
    vectors = np.vstack([
        np.frombuffer(row["embedding"], dtype=np.float32, count=row["embedding_dim"])
        for row in rows
    ])
    vector_scores = vectors @ query_vector
    keyword_scores = bm25_scores(question, [row["content"] for row in rows])
    candidate_indices = reciprocal_rank_fusion(vector_scores, keyword_scores, RETRIEVAL_CANDIDATES)
    candidates = [rows[index] for index in candidate_indices]
    return rerank(question, candidates, top_k)


def generate_answer(question: str, ranked: list[dict]) -> str:
    context_parts = []
    for number, item in enumerate(ranked, 1):
        row = item["row"]
        context_parts.append(
            f"[S{number}] 来源: {row['source']}，片段: {row['chunk_index']}\n{row['content']}"
        )
    context = "\n\n".join(context_parts)
    system = (
        "你是严谨的本地知识库助手。只能依据用户提供的资料片段回答。"
        "每个事实性结论后必须用 [S1] 这样的编号标注来源。"
        "如果资料不足、冲突或无法确定，明确说明，不得使用常识补全。"
        "先直接回答，再简要列出依据；不要暴露内部推理过程。"
    )
    user = f"问题：{question}\n\n资料片段：\n{context}"
    payload = {
        "model": LLM_MODEL,
        "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}],
        "temperature": 0.2,
        "top_p": 0.9,
        "max_tokens": 1200,
        "chat_template_kwargs": {"enable_thinking": False},
        "stream": False,
    }
    with httpx.Client(timeout=REQUEST_TIMEOUT, headers=headers('LLM')) as client:
        response = client.post(f"{LLM_URL}/chat/completions", json=payload)
        response.raise_for_status()
        return response.json()["choices"][0]["message"]["content"]


def iter_files(root: Path, recursive: bool) -> Iterable[Path]:
    if root.is_file():
        yield root
        return
    iterator = root.rglob("*") if recursive else root.glob("*")
    for path in iterator:
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS and path.resolve().is_relative_to(DOCUMENTS_DIR):
            yield path


@app.on_event("startup")
def initialize() -> None:
    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
    with connect_db():
        pass


@app.get("/health")
def health() -> dict:
    services = {}
    with httpx.Client(timeout=3) as client:
        for name, url in {"generator": LLM_URL, "embedding": EMBED_URL, "reranker": RERANK_URL}.items():
            try:
                kind = {'generator':'LLM','embedding':'EMBED','reranker':'RERANK'}[name]
                services[name] = client.get(f"{url}/models", headers=headers(kind)).status_code == 200
            except Exception:
                services[name] = False
    return {"ok": all(services.values()), "services": services}

@app.get('/ready')
def ready():
    return {'ok': True}

@app.get('/settings')
def settings():
    return {'documents_dir': str(DOCUMENTS_DIR)}


@app.post("/ingest")
def ingest(request: IngestRequest) -> dict:
    root = local_path(request.path)
    if not root.is_relative_to(DOCUMENTS_DIR):
        raise HTTPException(403, 'Import is restricted to RAG_DOCUMENTS_DIR. Move documents there first.')
    if not root.exists():
        raise HTTPException(status_code=404, detail=f"Path does not exist: {root}")
    files = list(iter_files(root, request.recursive))
    indexed_files = 0
    indexed_chunks = 0
    skipped_files = []
    with connect_db() as connection:
        for path in files:
            try:
                if not path.resolve().is_relative_to(DOCUMENTS_DIR):
                    raise ValueError('Document resolves outside allowed directory')
                if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                    raise ValueError('Unsupported document format')
                if path.stat().st_size > MAX_FILE_BYTES:
                    raise ValueError('Document exceeds configured size limit')
                payload = path.read_bytes()
                file_hash = hashlib.sha256(payload).hexdigest()
                existing = connection.execute(
                    "SELECT file_hash FROM chunks WHERE collection = ? AND source = ? LIMIT 1",
                    (request.collection, str(path)),
                ).fetchone()
                if existing and existing["file_hash"] == file_hash:
                    continue
                text = read_document(path)
                chunks = split_chunks(text)
                if not chunks:
                    skipped_files.append({"path": str(path), "reason": "no extractable text"})
                    continue
                title = document_title(text, path.stem)
                vectors = embed_texts(chunks)
                connection.execute(
                    "DELETE FROM chunks WHERE collection = ? AND source = ?",
                    (request.collection, str(path)),
                )
                connection.executemany(
                    """
                    INSERT INTO chunks
                    (collection, source, file_hash, chunk_index, title, content, embedding, embedding_dim)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            request.collection, str(path), file_hash, index, title, chunk,
                            vector.astype(np.float32).tobytes(), len(vector),
                        )
                        for index, (chunk, vector) in enumerate(zip(chunks, vectors), 1)
                    ],
                )
                connection.commit()
                indexed_files += 1
                indexed_chunks += len(chunks)
            except Exception as error:
                connection.rollback()
                skipped_files.append({"path": str(path), "reason": str(error)})
    return {
        "collection": request.collection,
        "found_files": len(files),
        "indexed_files": indexed_files,
        "indexed_chunks": indexed_chunks,
        "skipped": skipped_files,
    }


@app.post("/query")
def query(request: QueryRequest) -> dict:
    try:
        ranked = retrieve(request.question, request.collection, request.top_k)
        sources = [
            {
                "id": f"S{index}",
                "path": item["row"]["source"],
                "title": item["row"]["title"],
                "chunk": item["row"]["chunk_index"],
                "score": item["score"],
                "text": item["row"]["content"],
            }
            for index, item in enumerate(ranked, 1)
        ]
        answer = generate_answer(request.question, ranked) if request.generate else None
        return {"question": request.question, "answer": answer, "sources": sources}
    except httpx.HTTPError as error:
        raise HTTPException(status_code=502, detail=f"Local model service error: {error}") from error


@app.get("/collections")
def collections() -> dict:
    with connect_db() as connection:
        rows = connection.execute(
            "SELECT collection, COUNT(DISTINCT source) AS files, COUNT(*) AS chunks "
            "FROM chunks GROUP BY collection ORDER BY collection"
        ).fetchall()
    return {"collections": [dict(row) for row in rows]}
